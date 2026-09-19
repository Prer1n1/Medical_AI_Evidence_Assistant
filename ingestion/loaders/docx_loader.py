"""DOCX loader — one Document per heading-delimited section, plus one
Document per table. Ported verbatim from enterprise-agentic-rag: some
hospital policy/protocol documents are distributed as Word docs, and this
loader's structure (heading-section splitting, tables isolated so the
chunker can treat them row-wise instead of as prose) is entirely
format-generic, no medical-specific change needed.
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Union

import docx

from ..schema import Document, DocumentMetadata, canonical_source
from .base import BaseLoader


class DOCXLoader(BaseLoader):
    doc_type = "docx"

    def load(self, file_path: Union[str, Path]) -> List[Document]:
        file_path = Path(file_path)
        doc = docx.Document(str(file_path))

        props = doc.core_properties
        title = props.title or file_path.stem
        author = props.author or None
        created_date = str(props.created) if props.created else None

        documents: List[Document] = []
        current_section = "Document Start"
        buffer: List[str] = []

        def flush() -> None:
            text = "\n".join(buffer).strip()
            if text:
                documents.append(
                    Document(
                        content=text,
                        metadata=DocumentMetadata(
                            source=canonical_source(file_path),
                            doc_type=self.doc_type,
                            title=title,
                            author=author,
                            created_date=created_date,
                            section=current_section,
                        ),
                    )
                )

        for para in doc.paragraphs:
            style_name = (para.style.name or "").lower() if para.style else ""
            if style_name.startswith("heading"):
                flush()
                buffer = []
                current_section = para.text.strip() or current_section
            elif para.text.strip():
                buffer.append(para.text)
        flush()

        for table_index, table in enumerate(doc.tables, start=1):
            rows_text = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            table_content = "\n".join(r for r in rows_text if r.strip())
            if table_content:
                documents.append(
                    Document(
                        content=table_content,
                        metadata=DocumentMetadata(
                            source=canonical_source(file_path),
                            doc_type=self.doc_type,
                            title=title,
                            author=author,
                            created_date=created_date,
                            section=f"Table {table_index}",
                            extra={"is_table": True},
                        ),
                    )
                )

        return documents
